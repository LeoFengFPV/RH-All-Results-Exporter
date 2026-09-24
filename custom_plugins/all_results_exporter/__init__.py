"""Export saved RotorHazard races to one XLSX sheet. Reads data on demand."""

import io
import logging
from collections import defaultdict

from eventmanager import Evt
from data_export import DataExporter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
ALT_FILL = PatternFill("solid", fgColor="F2F6FC")
CENTER = Alignment(horizontal="center", vertical="center")


def _get(obj, key, default=None):
    return obj.get(key, default) if isinstance(obj, dict) else getattr(obj, key, default)


def _id(value):
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return str(value)


def _sort_number(value):
    value = _id(value)
    return (0, value) if isinstance(value, int) else (1, str(value))


def _board(result):
    """Both a direct leaderboard and a wrapped leaderboard are accepted."""
    if not isinstance(result, dict):
        return [], {}
    board = result.get("leaderboard", result)
    if not isinstance(board, dict):
        return [], {}
    meta = board.get("meta") or {}
    primary = meta.get("primary_leaderboard")
    for key in ([primary] if primary else []) + [
            "by_race_time", "by_fastest_lap", "by_consecutives"]:
        if key and isinstance(board.get(key), list):
            return board[key], meta
    return [], meta


def _laps(records):
    """Deleted crossings do not appear in the public per-lap results."""
    ordered = sorted(records, key=lambda lap: (
        _get(lap, "lap_time_stamp", 0) or 0, _sort_number(_get(lap, "id"))))
    return [str(_get(lap, "lap_time_formatted")) for lap in ordered
            if not _get(lap, "deleted") and _get(lap, "lap_time_formatted")]


def assemble_all_results(rhapi):
    """Join by saved race / pilot-run IDs, never by duplicate callsigns."""
    try:
        db = rhapi.db
        races = list(db.races)
        if not races:
            logger.warning("No saved races to export")
            return None

        heats = {_id(_get(h, "id")): h for h in db.heats}
        classes = {_id(_get(c, "id")): c for c in db.raceclasses}
        pilots = {_id(_get(p, "id")): p for p in db.pilots}
        formats = {_id(_get(f, "id")): f for f in db.raceformats}
        runs_by_race = defaultdict(list)
        for run in db.pilotruns:
            runs_by_race[_id(_get(run, "race_id"))].append(run)
        laps_by_run = defaultdict(list)
        for lap in db.laps:
            laps_by_run[_id(_get(lap, "pilotrace_id"))].append(lap)

        rows, lap_numbers = [], set()
        races.sort(key=lambda race: (
            _sort_number(_get(race, "heat_id")),
            _sort_number(_get(race, "round_id")),
            _sort_number(_get(race, "id"))))
        for race in races:
            race_id = _id(_get(race, "id"))
            heat_id = _id(_get(race, "heat_id"))
            class_id = _id(_get(race, "class_id"))
            result, meta = _board(db.race_results(race_id))
            fmt = formats.get(_id(_get(race, "format_id")))
            start_behavior = meta.get("start_behavior", _get(fmt, "start_behavior"))
            try:
                start_behavior = int(start_behavior)
            except (TypeError, ValueError):
                start_behavior = None
            first_lap = 1 if start_behavior == 1 else 0
            total_key = "total_time_laps" if start_behavior == 2 else "total_time"

            by_pilot, by_node = {}, {}
            for entry in result:
                pilot_key = _id(entry.get("pilot_id"))
                node_key = _id(entry.get("node", entry.get("node_index")))
                if pilot_key is not None:
                    if pilot_key in by_pilot:
                        raise ValueError("Duplicate leaderboard pilot ID in race %s" % race_id)
                    by_pilot[pilot_key] = entry
                if node_key is not None:
                    if node_key in by_node:
                        raise ValueError("Duplicate leaderboard seat in race %s" % race_id)
                    by_node[node_key] = entry

            runs = sorted(runs_by_race.get(race_id, ()),
                          key=lambda x: _sort_number(_get(x, "node_index")))
            if not runs:
                logger.warning("Saved race %s has no pilot runs", race_id)
            for run in runs or [None]:
                pilot_id = _id(_get(run, "pilot_id"))
                # An unassigned seat (a real run with no pilot) is not a result; skip it.
                # A missing run (whole race has no runs) is kept below as a placeholder.
                if run is not None and pilot_id is None:
                    continue
                node = _id(_get(run, "node_index"))
                entry = by_pilot.get(pilot_id)
                if entry is None and node is not None:
                    candidate = by_node.get(node)
                    if candidate and _id(candidate.get("pilot_id")) in (None, pilot_id):
                        entry = candidate
                pilot = pilots.get(pilot_id)
                values = _laps(laps_by_run.get(_id(_get(run, "id")), ()))
                lap_times = {first_lap + i: val for i, val in enumerate(values)}
                lap_numbers.update(lap_times)
                note = ("无飞手记录" if run is None else
                        "未匹配排行榜，请核对名次及圈数" if entry is None else "")
                if entry is not None and start_behavior is None:
                    note = "起跑方式未知，请核对圈号及总时间"
                if entry is not None and entry.get(total_key) is None:
                    note = (note + "；" if note else "") + "总时间字段缺失，请核对"
                if run is not None and entry is None:
                    logger.warning("Race %s run %s has no matching leaderboard entry",
                                   race_id, _get(run, "id"))
                rows.append({
                    "class_name": str(_get(classes.get(class_id), "name") or ""),
                    "heat_name": str(_get(heats.get(heat_id), "name") or heat_id or ""),
                    "round_id": _get(race, "round_id", ""),
                    "position": entry.get("position", "") if entry else "",
                    "callsign": str(_get(pilot, "callsign") or
                                    (entry.get("callsign") if entry else "") or ""),
                    "team": str((entry.get("team_name") if entry else None)
                                or _get(pilot, "team") or ""),
                    "frequency": (_get(run, "frequency") if
                                  _get(run, "frequency") not in (None, 0) else ""),
                    "laps": entry.get("laps", "") if entry else "",
                    "average_lap": entry.get("average_lap", "") if entry else "",
                    "fastest_lap": entry.get("fastest_lap", "") if entry else "",
                    "lap_times": lap_times,
                    "total_time": entry.get(total_key, "") if entry else "",
                    "note": note,
                    "_class_id": class_id, "_heat_id": heat_id, "_race_id": race_id,
                })
        return {"rows": rows, "lap_numbers": sorted(lap_numbers),
                "race_count": len(races)}
    except Exception:
        logger.exception("Could not assemble all saved race results")
        return None


def _merge_groups(ws, rows, col, level):
    """Only merge within an identical class / heat / saved race ID prefix."""
    keys = ("_class_id", "_heat_id", "_race_id")[:level]
    begin = 0
    while begin < len(rows):
        end = begin + 1
        group = tuple(rows[begin][key] for key in keys)
        while end < len(rows) and tuple(rows[end][key] for key in keys) == group:
            end += 1
        if end - begin > 1:
            ws.merge_cells(start_row=begin + 2, start_column=col,
                           end_row=end + 1, end_column=col)
            ws.cell(begin + 2, col).alignment = CENTER
        begin = end


def write_excel(payload):
    if not payload or not payload.get("rows"):
        return None
    rows = payload["rows"]
    columns = [
        ("class_name", "类别", 14), ("heat_name", "分组", 12),
        ("round_id", "轮次", 7), ("position", "名次", 7),
        ("callsign", "飞手", 16), ("team", "团队", 12),
        ("frequency", "频点 (MHz)", 12), ("laps", "圈数", 7),
        ("average_lap", "平均圈速", 12), ("fastest_lap", "最快圈速", 12),
        ("note", "备注", 28),
    ]
    columns.extend(("lap_%s" % n, "第%s圈" % n, 12)
                   for n in payload["lap_numbers"])
    columns.append(("total_time", "总时间", 14))
    active = [(key, title, width) for key, title, width in columns
              if key.startswith("lap_") or any(
                  row.get(key) not in (None, "") for row in rows)]
    wb = Workbook()
    ws = wb.active
    ws.title = "比赛结果"
    ws.append([title for _, title, _ in active])
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    for col, (_, _, width) in enumerate(active, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(1, col)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, CENTER

    for idx, row in enumerate(rows, 2):
        ws.append([row["lap_times"].get(int(key[4:]), "")
                   if key.startswith("lap_") else row.get(key, "")
                   for key, _, _ in active])
        # ws[idx] recomputes max_column from all existing cells on each row.
        for col in range(1, len(active) + 1):
            cell = ws.cell(idx, col)
            if idx % 2 == 0:
                cell.fill = ALT_FILL
            cell.alignment = CENTER
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

    for level, key in enumerate(("class_name", "heat_name", "round_id"), 1):
        col = next((i for i, (name, _, _) in enumerate(active, 1)
                    if name == key), None)
        if col:
            _merge_groups(ws, rows, col, level)
    try:
        with io.BytesIO() as buf:
            wb.save(buf)
            return {
                "data": buf.getvalue(),
                "encoding": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "ext": "xlsx",
            }
    except Exception:
        logger.exception("Failed to generate XLSX")
        return None
    finally:
        wb.close()


def register_handlers(args):
    args["register_fn"](DataExporter(
        "Export All Results (XLSX)", write_excel, assemble_all_results,
        name="all_results_xlsx"))


def initialize(rhapi):
    rhapi.events.on(Evt.DATA_EXPORT_INITIALIZE, register_handlers)
