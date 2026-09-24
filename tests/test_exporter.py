"""Regression tests for the All Results Exporter plugin.

All data is simulated; no real RotorHazard database is read. The RotorHazard
modules the plugin imports are stubbed before the plugin is loaded.
"""

import importlib.util
import io
import os
import sys
import types

# --- Stub RotorHazard modules that the plugin imports ----------------------
if "eventmanager" not in sys.modules:
    _ev = types.ModuleType("eventmanager")

    class Evt:  # minimal stand-in
        DATA_EXPORT_INITIALIZE = "data_export_initialize"

    _ev.Evt = Evt
    sys.modules["eventmanager"] = _ev

if "data_export" not in sys.modules:
    _de = types.ModuleType("data_export")

    class DataExporter:
        def __init__(self, *args, **kwargs):
            pass

    _de.DataExporter = DataExporter
    sys.modules["data_export"] = _de

# --- Load the plugin module by file path -----------------------------------
_HERE = os.path.dirname(os.path.abspath(__file__))
_PLUGIN_PATH = os.path.join(
    _HERE, "..", "custom_plugins", "all_results_exporter", "__init__.py")
_spec = importlib.util.spec_from_file_location("all_results_exporter", _PLUGIN_PATH)
exporter = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(exporter)

from openpyxl import load_workbook  # noqa: E402


# --- helpers ---------------------------------------------------------------
def fmt(ms):
    seconds = ms / 1000
    minutes = int(seconds // 60)
    return "%d:%06.3f" % (minutes, seconds - minutes * 60)


class MockDB:
    def __init__(self):
        self.races_ = []
        self.heats_ = []
        self.classes_ = []
        self.pilots_ = []
        self.formats_ = []
        self.runs_ = []
        self.laps_ = []
        self.results_ = {}

    @property
    def races(self):
        return list(self.races_)

    @property
    def heats(self):
        return list(self.heats_)

    @property
    def raceclasses(self):
        return list(self.classes_)

    @property
    def pilots(self):
        return list(self.pilots_)

    @property
    def raceformats(self):
        return list(self.formats_)

    @property
    def pilotruns(self):
        return list(self.runs_)

    @property
    def laps(self):
        return list(self.laps_)

    def race_results(self, race_id):
        return self.results_.get(race_id, {"leaderboard": {}})


class MockRHAPI:
    def __init__(self):
        self.db = MockDB()


def add_format(db, fid, start_behavior):
    db.formats_.append({"id": fid, "start_behavior": start_behavior})


def add_class(db, cid, name):
    db.classes_.append({"id": cid, "name": name})


def add_heat(db, hid, name):
    db.heats_.append({"id": hid, "name": name})


def add_pilot(db, pid, callsign, team="", name=None):
    db.pilots_.append(
        {"id": pid, "callsign": callsign, "team": team, "name": name or callsign})


def add_race(db, rid, hid, cid, fid, round_id=1):
    db.races_.append({"id": rid, "heat_id": hid, "class_id": cid,
                      "format_id": fid, "round_id": round_id})


def add_run(db, run_id, rid, node, pilot, freq):
    db.runs_.append({"id": run_id, "race_id": rid, "node_index": node,
                     "pilot_id": pilot, "frequency": freq})


def add_laps(db, run_id, lap_ms, deleted_indexes=()):
    timestamp = 0
    for i, ms in enumerate(lap_ms):
        timestamp += ms
        db.laps_.append({
            "id": len(db.laps_) + 1, "pilotrace_id": run_id,
            "lap_time_stamp": timestamp, "lap_time": ms,
            "lap_time_formatted": fmt(ms),
            "deleted": 1 if i in deleted_indexes else 0})


def leaderboard_entry(pid, node, position, laps, total, total_laps,
                      callsign=None, team="A", avg="0:07.000", fast="0:07.000"):
    return {"pilot_id": pid, "node": node, "position": position, "laps": laps,
            "total_time": total, "total_time_laps": total_laps,
            "average_lap": avg, "fastest_lap": fast,
            "callsign": callsign, "team_name": team}


def set_leaderboard(db, rid, start_behavior, entries):
    db.results_[rid] = {"leaderboard": {
        "meta": {"primary_leaderboard": "by_race_time",
                 "start_behavior": start_behavior},
        "by_race_time": entries}}


def assemble(api):
    payload = exporter.assemble_all_results(api)
    assert payload is not None
    return payload


def render_sheet(payload):
    result = exporter.write_excel(payload)
    assert result is not None
    wb = load_workbook(io.BytesIO(result["data"]))
    return wb, wb.active


def headers_of(ws):
    return [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]


# --- tests -----------------------------------------------------------------
def test_1_duplicate_callsign_no_cross_talk():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 0)
    add_class(db, 1, "C")
    add_heat(db, 1, "H")
    add_race(db, 1, 1, 1, 1)
    add_pilot(db, 10, "Same")
    add_pilot(db, 11, "Same")
    add_run(db, 100, 1, 0, 10, 5800)
    add_run(db, 101, 1, 1, 11, 5820)
    add_laps(db, 100, [8000, 7000, 7000])
    add_laps(db, 101, [9000, 8000, 8000])
    set_leaderboard(db, 1, 0, [
        leaderboard_entry(10, 0, 1, 2, "0:22.000", "0:14.000", callsign="Same"),
        leaderboard_entry(11, 1, 2, 2, "0:25.000", "0:16.000", callsign="Same")])

    rows = assemble(api)["rows"]
    assert len(rows) == 2
    r10 = next(r for r in rows if r["frequency"] == 5800)
    r11 = next(r for r in rows if r["frequency"] == 5820)
    assert r10["lap_times"][0] == "0:08.000"
    assert r11["lap_times"][0] == "0:09.000"
    assert r10["lap_times"][1] == "0:07.000"
    assert r11["lap_times"][1] == "0:08.000"
    assert r10["callsign"] == "Same" and r11["callsign"] == "Same"


def test_2_empty_seat_filtered():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 0)
    add_class(db, 1, "C")
    add_heat(db, 1, "H")
    add_race(db, 1, 1, 1, 1)
    add_pilot(db, 10, "Pilot A")
    add_run(db, 100, 1, 0, None, 5658)   # unassigned seat
    add_run(db, 101, 1, 1, 10, 5800)
    add_laps(db, 101, [8000, 7000])
    set_leaderboard(db, 1, 0, [
        leaderboard_entry(10, 1, 1, 1, "0:15.000", "0:07.000", callsign="Pilot A")])

    rows = assemble(api)["rows"]
    assert len(rows) == 1
    assert rows[0]["frequency"] == 5800
    assert rows[0]["callsign"] == "Pilot A"


def test_3_race_without_runs_keeps_placeholder():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 0)
    add_class(db, 1, "C")
    add_heat(db, 1, "H")
    add_race(db, 1, 1, 1, 1)
    set_leaderboard(db, 1, 0, [])

    rows = assemble(api)["rows"]
    assert len(rows) == 1
    assert rows[0]["note"] == "无飞手记录"


def test_4_deleted_laps_excluded():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 0)
    add_class(db, 1, "C")
    add_heat(db, 1, "H")
    add_race(db, 1, 1, 1, 1)
    add_pilot(db, 10, "Pilot A")
    add_run(db, 100, 1, 0, 10, 5800)
    add_laps(db, 100, [8000, 7000, 7000], deleted_indexes=(1,))
    set_leaderboard(db, 1, 0, [
        leaderboard_entry(10, 0, 1, 1, "0:15.000", "0:07.000", callsign="Pilot A")])

    lap_times = assemble(api)["rows"][0]["lap_times"]
    assert set(lap_times.keys()) == {0, 1}
    assert lap_times[0] == "0:08.000"
    assert lap_times[1] == "0:07.000"


def test_5_holeshot_starts_at_lap0():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 0)
    add_class(db, 1, "C")
    add_heat(db, 1, "H")
    add_race(db, 1, 1, 1, 1)
    add_pilot(db, 10, "Pilot A")
    add_run(db, 100, 1, 0, 10, 5800)
    add_laps(db, 100, [8000, 7000, 7000])
    set_leaderboard(db, 1, 0, [
        leaderboard_entry(10, 0, 1, 2, "0:22.000", "0:14.000", callsign="Pilot A")])

    lap_times = assemble(api)["rows"][0]["lap_times"]
    assert min(lap_times.keys()) == 0
    assert lap_times[0] == "0:08.000"


def test_6_firstlap_starts_at_lap1():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 1)
    add_class(db, 1, "C")
    add_heat(db, 1, "H")
    add_race(db, 1, 1, 1, 1)
    add_pilot(db, 10, "Pilot A")
    add_run(db, 100, 1, 0, 10, 5800)
    add_laps(db, 100, [8000, 7000, 7000])
    set_leaderboard(db, 1, 1, [
        leaderboard_entry(10, 0, 1, 3, "0:22.000", "0:22.000", callsign="Pilot A")])

    lap_times = assemble(api)["rows"][0]["lap_times"]
    assert 0 not in lap_times
    assert set(lap_times.keys()) == {1, 2, 3}
    assert lap_times[1] == "0:08.000"


def test_7_staggered_uses_total_time_laps():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 2)
    add_class(db, 1, "C")
    add_heat(db, 1, "H")
    add_race(db, 1, 1, 1, 1)
    add_pilot(db, 10, "Pilot A")
    add_run(db, 100, 1, 0, 10, 5800)
    add_laps(db, 100, [8000, 7000, 7000])
    set_leaderboard(db, 1, 2, [
        leaderboard_entry(10, 0, 1, 2, "0:99.999", "0:14.000", callsign="Pilot A")])

    row = assemble(api)["rows"][0]
    assert row["total_time"] == "0:14.000"
    assert min(row["lap_times"].keys()) == 0


def test_8_same_heat_name_different_class_not_merged():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 0)
    add_class(db, 1, "Class A")
    add_class(db, 2, "Class B")
    add_heat(db, 1, "X")
    add_heat(db, 2, "X")
    add_race(db, 1, 1, 1, 1)
    add_race(db, 2, 2, 2, 1)
    add_pilot(db, 10, "Pilot A")
    add_pilot(db, 11, "Pilot B")
    add_pilot(db, 12, "Pilot C")
    add_pilot(db, 13, "Pilot D")
    add_run(db, 100, 1, 0, 10, 5800)
    add_run(db, 101, 1, 1, 11, 5820)
    add_run(db, 102, 2, 0, 12, 5840)
    add_run(db, 103, 2, 1, 13, 5860)
    for run_id in (100, 101, 102, 103):
        add_laps(db, run_id, [8000, 7000])
    set_leaderboard(db, 1, 0, [
        leaderboard_entry(10, 0, 1, 1, "0:15.000", "0:07.000", callsign="Pilot A"),
        leaderboard_entry(11, 1, 2, 1, "0:15.000", "0:07.000", callsign="Pilot B")])
    set_leaderboard(db, 2, 0, [
        leaderboard_entry(12, 0, 1, 1, "0:15.000", "0:07.000", callsign="Pilot C"),
        leaderboard_entry(13, 1, 2, 1, "0:15.000", "0:07.000", callsign="Pilot D")])

    _, ws = render_sheet(assemble(api))
    heat_col = headers_of(ws).index("分组") + 1
    heat_merges = sorted(
        str(m) for m in ws.merged_cells.ranges
        if m.min_col == heat_col and m.max_col == heat_col)
    assert "B2:B3" in heat_merges
    assert "B4:B5" in heat_merges
    # no merge spans the boundary between the two real races
    assert not any(m.min_row == 2 and m.max_row == 5
                   for m in ws.merged_cells.ranges if m.min_col == heat_col)


def test_9_formula_like_callsign_stored_as_text():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 0)
    add_class(db, 1, "C")
    add_heat(db, 1, "H")
    add_race(db, 1, 1, 1, 1)
    add_pilot(db, 10, "=1+1")
    add_run(db, 100, 1, 0, 10, 5800)
    add_laps(db, 100, [8000, 7000])
    set_leaderboard(db, 1, 0, [
        leaderboard_entry(10, 0, 1, 1, "0:15.000", "0:07.000", callsign="=1+1")])

    _, ws = render_sheet(assemble(api))
    callsign_col = headers_of(ws).index("飞手") + 1
    cell = ws.cell(2, callsign_col)
    assert cell.data_type == "s"
    assert cell.value == "=1+1"


def test_10_total_time_is_last_column():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 0)
    add_class(db, 1, "C")
    add_heat(db, 1, "H")
    add_race(db, 1, 1, 1, 1)
    add_pilot(db, 10, "Pilot A")
    add_run(db, 100, 1, 0, 10, 5800)
    add_laps(db, 100, [8000, 7000])
    set_leaderboard(db, 1, 0, [
        leaderboard_entry(10, 0, 1, 1, "0:15.000", "0:07.000", callsign="Pilot A")])

    _, ws = render_sheet(assemble(api))
    assert headers_of(ws)[-1] == "总时间"


def test_11_workbook_reopens_with_openpyxl():
    api = MockRHAPI()
    db = api.db
    add_format(db, 1, 0)
    add_class(db, 1, "C")
    add_heat(db, 1, "H")
    add_race(db, 1, 1, 1, 1)
    add_pilot(db, 10, "Pilot A")
    add_run(db, 100, 1, 0, 10, 5800)
    add_laps(db, 100, [8000, 7000])
    set_leaderboard(db, 1, 0, [
        leaderboard_entry(10, 0, 1, 1, "0:15.000", "0:07.000", callsign="Pilot A")])

    result = exporter.write_excel(assemble(api))
    reopened = load_workbook(io.BytesIO(result["data"]))
    assert reopened.active.max_row >= 2
